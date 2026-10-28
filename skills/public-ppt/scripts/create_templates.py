#!/usr/bin/env python3
"""Rebuild the original, editable business and Internet briefing template decks."""
from pathlib import Path
import json

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
FONT = "PingFang SC"


def text(slide, name, value, x, y, w, h, size=21, color="22334A", bold=False):
    """Add a native editable text box with explicit Chinese and Latin fonts."""
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.025)
    frame.margin_top = frame.margin_bottom = Inches(0.025)
    for index, line in enumerate(value.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(10)
        run = paragraph.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)
        east_asian = OxmlElement("a:ea")
        east_asian.set("typeface", FONT)
        run._r.get_or_add_rPr().append(east_asian)
    return shape


def table(slide, name, rows, x=0.8, y=2.1, w=11.7, h=3.65, accent="153B65"):
    """Add an editable native table; row heights remain uniform."""
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    for ri, row in enumerate(rows):
        for ci, value in enumerate(row):
            cell = shape.table.cell(ri, ci)
            cell.text = value
            cell.margin_left = Inches(0.15)
            cell.margin_top = Inches(0.11)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(accent if ri == 0 else ("F1F4F7" if ri % 2 else "FFFFFF"))
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.name = FONT
                    run.font.size = Pt(18)
                    run.font.bold = ri == 0
                    run.font.color.rgb = RGBColor.from_string("FFFFFF" if ri == 0 else "22334A")
                    ea = OxmlElement("a:ea")
                    ea.set("typeface", FONT)
                    run._r.get_or_add_rPr().append(ea)
    return shape


def create(kind):
    """Create one complete twelve-slide briefing and an authored example spec."""
    business = kind == "business"
    accent = "153B65" if business else "3855D9"
    dark = "10283E" if business else "131B35"
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    values = {}
    chart_edits = []
    def token(key, value):
        values["{{" + key + "}}"] = value
        return "{{" + key + "}}"
    def slide(title, subtitle="", cover=False):
        item = presentation.slides.add_slide(presentation.slide_layouts[6])
        item.background.fill.solid()
        item.background.fill.fore_color.rgb = RGBColor.from_string(dark if cover else "FFFFFF")
        page = len(presentation.slides)
        text(item, "page-number", f"{page:02}", 12, 6.92, 0.5, 0.28, 11, "A9B8CC" if cover else "52637B")
        text(item, f"s{page}-title", title, 0.8, 1.7 if cover else 0.65, 11.7, 1.6 if cover else 0.8, 42 if cover else 32, "FFFFFF" if cover else dark, True)
        if subtitle:
            text(item, f"s{page}-subtitle", subtitle, 0.8, 3.55 if cover else 1.5, 11.5, 0.5, 19, "CED8E9" if cover else "52637B")
        return item
    def note(item, content):
        item.notes_slide.notes_text_frame.text = content
    def foot(item, key="source", value="来源：用户材料，具体页码待补充"):
        text(item, f"s{len(presentation.slides)}-source", token(f"s{len(presentation.slides)}_{key}", value), 0.8, 6.65, 10.8, 0.4, 11, "52637B")
    def rows(item, entries):
        for index, (heading, body) in enumerate(entries):
            y = 2.15 + index * 1.28
            text(item, f"s{len(presentation.slides)}-heading-{index}", heading, 0.8, y, 2.45, 0.6, 23, accent, True)
            text(item, f"s{len(presentation.slides)}-body-{index}", body, 3.5, y, 8.7, 0.9, 21)

    item = slide(token("cover_title", "年度经营工作汇报" if business else "产品与增长复盘"), token("cover_subtitle", "报告期、汇报单位与汇报人待补充"), True)
    text(item, "cover-topic", "经营分析 · 重点项目 · 资源决策" if business else "用户价值 · 核心指标 · 实验结论", 0.8, 5.55, 11, 0.6, 21, "CED8E9")
    note(item, "原生可编辑封面。标题、报告期、汇报人依据用户材料填写。")

    item = slide("汇报议程" if business else "本次需要讨论的事项")
    rows(item, [("01", token("agenda_1", "经营结果与关键变化" if business else "目标达成与核心指标")), ("02", token("agenda_2", "重点任务、风险与资源安排" if business else "用户反馈、漏斗与实验结果")), ("03", token("agenda_3", "下一步工作及提请决策事项" if business else "产品路线图与资源决策"))])

    item = slide(token("s3_title", "关键工作结论" if business else "目标达成情况"))
    rows(item, [(token("s3_h1", "工作结果" if business else "目标结果"), token("s3_b1", "主要结果待补充，并注明对照目标与统计口径。")), (token("s3_h2", "关键原因" if business else "用户变化"), token("s3_b2", "依据材料填写关键变化及其证据。")), (token("s3_h3", "后续动作"), token("s3_b3", "填写责任人、完成时间和验收条件。"))])
    foot(item)

    item = slide("核心指标" if business else "北极星指标与护栏指标", token("s4_subtitle", "指标口径、时间范围与数据来源待补充"))
    for index, (name, value, detail) in enumerate([( "营收" if business else "核心活跃用户", "待补充", "目标与同比待补充"), ("利润" if business else "有效转化率", "待补充", "分母与样本量待补充"), ("回款" if business else "次月留存率", "待补充", "统计窗口待补充")]):
        x = 0.8 + index * 4.15
        text(item, f"metric-{index}-value", token(f"metric_{index}_value", value), x, 2.65, 3.7, 0.95, 40, accent, True)
        text(item, f"metric-{index}-name", token(f"metric_{index}_name", name), x, 3.8, 3.7, 0.6, 24, dark, True)
        text(item, f"metric-{index}-detail", token(f"metric_{index}_detail", detail), x, 4.65, 3.7, 1.05, 19)
    foot(item)

    item = slide("经营趋势" if business else "关键转化环节", token("s5_subtitle", "模板演示数据：交付前用已核实数据替换并核对单位"))
    data = CategoryChartData()
    data.categories = ["第一期", "第二期", "第三期", "第四期"] if business else ["访问", "注册", "激活", "付费"]
    data.add_series("演示数据", [30, 45, 60, 75] if business else [100, 60, 35, 12])
    graphic = item.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED if business else XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.85), Inches(2.3), Inches(11.6), Inches(3.95), data)
    graphic.name = "editable-metric-chart"
    chart = graphic.chart
    chart.has_legend = False
    chart.chart_style = 10
    chart.category_axis.tick_labels.font.name = FONT
    chart.category_axis.tick_labels.font.size = Pt(17)
    chart.value_axis.tick_labels.font.name = FONT
    chart.value_axis.tick_labels.font.size = Pt(14)
    chart.series[0].format.fill.solid()
    chart.series[0].format.fill.fore_color.rgb = RGBColor.from_string(accent)
    foot(item, value="演示数据：30/45/60/75" if business else "演示数据：100/60/35/12；各环节分母待定义")
    chart_edits.append({"slide": 5, "shape_id": graphic.shape_id, "categories": (["第一期", "第二期", "第三期", "第四期"] if business else ["访问", "注册", "激活", "付费"]), "series": [{"name": "演示数据", "values": [30, 45, 60, 75] if business else [100, 60, 35, 12]}]})
    note(item, "图表和嵌入工作簿均可编辑。示例数字仅用于展示布局，正式交付必须替换并注明来源。")

    item = slide("重点项目进展" if business else "用户分层与留存")
    headers = ["重点项目", "当前进展", "责任部门", "验收时间"] if business else ["用户分层", "样本量", "次月留存", "主要问题"]
    data_rows = [headers]
    for ri in range(3):
        defaults = [f"项目{ri+1}", "进展待补充", "责任待补充", "时间待补充"] if business else [f"用户群{ri+1}", "待补充", "待补充", "反馈待补充"]
        data_rows.append([token(f"s6_r{ri}_c{ci}", value) for ci, value in enumerate(defaults)])
    table(item, "progress-table" if business else "cohort-table", data_rows, accent=accent)
    foot(item)

    item = slide("重点任务与实施措施" if business else "实验记录与结论")
    if business:
        rows(item, [(token(f"s7_h{i}", h), token(f"s7_b{i}", b)) for i, (h, b) in enumerate([( "任务一", "工作内容、责任分工、完成标准待补充。"), ("任务二", "资源需求、协同部门、检查节点待补充。"), ("任务三", "当前障碍、解决措施、预期完成日期待补充。")])])
    else:
        table(item, "experiment-table", [["实验", "样本与周期", "结果", "后续动作"]] + [[token(f"s7_r{r}_c{c}", val) for c, val in enumerate([f"实验{r+1}", "待补充", "待验证", "待补充"])] for r in range(3)], accent=accent)
    foot(item)

    item = slide("风险与应对" if business else "产品路线图")
    headers = ["主要风险", "影响", "应对措施", "责任人"] if business else ["里程碑", "交付内容", "负责人", "验收标准"]
    table(item, "risk-table" if business else "roadmap-table", [headers] + [[token(f"s8_r{r}_c{c}", val) for c, val in enumerate([f"事项{r+1}", "待补充", "待补充", "待补充"])] for r in range(3)], accent=accent)
    foot(item)

    item = slide("资源与预算安排" if business else "产品与技术依赖")
    headers = ["资源项", "现有安排", "新增需求", "使用条件"] if business else ["依赖事项", "负责团队", "交付节点", "阻塞影响"]
    table(item, "resource-table" if business else "dependency-table", [headers] + [[token(f"s9_r{r}_c{c}", val) for c, val in enumerate([f"事项{r+1}", "待补充", "待补充", "待补充"])] for r in range(3)], accent=accent)
    foot(item)

    item = slide("工作计划与节点" if business else "风险与质量护栏")
    headers = ["阶段", "完成事项", "责任人", "检查节点"] if business else ["风险", "监测指标", "触发条件", "处理责任"]
    table(item, "timeline-table" if business else "guardrail-table", [headers] + [[token(f"s10_r{r}_c{c}", val) for c, val in enumerate([f"阶段{r+1}" if business else f"风险{r+1}", "待补充", "待补充", "待补充"])] for r in range(3)], accent=accent)
    foot(item)

    item = slide("提请研究决定事项" if business else "资源与决策请求")
    rows(item, [(token("decision_1_heading", "决策事项"), token("decision_1_body", "填写需要决定的具体事项与建议选项。")), (token("decision_2_heading", "依据与影响"), token("decision_2_body", "填写成本、收益、风险及其材料依据。")), (token("decision_3_heading", "办理要求"), token("decision_3_body", "填写决策期限、责任主体与后续办理安排。"))])
    foot(item)

    item = slide(token("closing_title", "请予审议" if business else "讨论与决策"), token("closing_subtitle", "后续办理人与联系方式待补充"), True)
    text(item, "closing-request", token("closing_request", "请就材料中的具体决策事项提出意见。"), 0.8, 5.1, 11.3, 1, 24, "CED8E9")
    directory = ROOT / ("public-ppt-business" if business else "public-ppt-internet") / "templates"
    directory.mkdir(parents=True, exist_ok=True)
    presentation.core_properties.title = "国企商务汇报原创模板" if business else "互联网产品与增长汇报原创模板"
    presentation.core_properties.author = "TianCode"
    presentation.core_properties.subject = "Original editable slide layouts; example figures are demonstration data"
    presentation.save(directory / "template.pptx")
    (directory / "example.json").write_text(json.dumps({"replacements": values, "charts": chart_edits}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(directory / "template.pptx")


if __name__ == "__main__":
    create("business")
    create("internet")
