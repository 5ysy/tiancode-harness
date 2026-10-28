#!/usr/bin/env python3
"""Exercise atomic edits, formatting retention, strict addressing, and chart data."""
import json
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

import deck
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[2]


class DeckWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)
        self.source = ROOT / "public-ppt-business/templates/template.pptx"
        self.spec = self.directory / "content.json"
        self.output = self.directory / "output.pptx"

    def write(self, value):
        self.spec.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")

    def test_text_edit_retains_untouched_parts_and_run_styles(self):
        self.write({"replacements": {"{{cover_title}}": "工作情况汇报"}})
        deck.fill(self.source, self.spec, self.output)
        with ZipFile(self.source) as source, ZipFile(self.output) as output:
            for name in source.namelist():
                if name != "ppt/slides/slide1.xml":
                    self.assertEqual(source.read(name), output.read(name), name)
            left = deck.xml(source.read("ppt/slides/slide1.xml"))
            right = deck.xml(output.read("ppt/slides/slide1.xml"))
            self.assertEqual(left.xpath("//a:rPr/@sz", namespaces=deck.NS), right.xpath("//a:rPr/@sz", namespaces=deck.NS))
        self.assertIn("工作情况汇报", str(deck.inspect(self.output)))

    def test_wrong_expected_text_leaves_no_output(self):
        self.write({"edits": [{"slide": 1, "shape_id": 3, "paragraph": 0, "expected_text": "outdated title", "text": "new title"}]})
        with self.assertRaises(ValueError):
            deck.fill(self.source, self.spec, self.output)
        self.assertFalse(self.output.exists())

    def test_invalid_page_and_missing_token_are_rejected(self):
        for spec in ({"selected_slides": [0]}, {"selected_slides": [1, 1]}, {"replacements": {"missing-token": "value"}}, {"unexpected": "value"}):
            self.write(spec)
            with self.assertRaises(ValueError):
                deck.fill(self.source, self.spec, self.output)
            self.assertFalse(self.output.exists())

    def test_category_chart_and_workbook_update_together(self):
        value = json.loads((ROOT / "public-ppt-business/templates/example.json").read_text())
        value["charts"][0]["categories"] = ["一期", "二期"]
        value["charts"][0]["series"] = [{"name": "核实数据", "values": [12, 18]}]
        self.write(value)
        deck.fill(self.source, self.spec, self.output)
        slide = Presentation(self.output).slides[4]
        chart = next(shape.chart for shape in slide.shapes if shape.has_chart)
        self.assertEqual(list(chart.series[0].values), [12.0, 18.0])
        from io import BytesIO
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(chart.part.chart_workbook.xlsx_part.blob), data_only=True)
        self.assertEqual(workbook.active.cell(3, 2).value, 18)
        self.assertEqual(deck.validate(self.output)["errors"], [])

    def test_nonfinite_chart_values_leave_no_output(self):
        value = json.loads((ROOT / "public-ppt-business/templates/example.json").read_text())
        value["charts"][0]["series"][0]["values"][0] = float("nan")
        self.write(value)
        with self.assertRaises(ValueError):
            deck.fill(self.source, self.spec, self.output)
        self.assertFalse(self.output.exists())

    def test_selection_changes_order_and_preserves_source(self):
        self.write({"selected_slides": [12, 1]})
        deck.fill(self.source, self.spec, self.output)
        result = deck.inspect(self.output)
        self.assertEqual(result["slide_count"], 2)
        self.assertIn("{{closing_title}}", str(result["slides"][0]))
        self.assertEqual(deck.inspect(self.source)["slide_count"], 12)


if __name__ == "__main__":
    unittest.main()
